#!/usr/bin/env python3
"""生成测试用的 Excel 文件（尽量覆盖各种真实结构）。

用法: python3 test/make-fixtures.py
依赖: openpyxl（本机已装）；xlsxwriter 可选，用于生成「已有背景水印」的文件
"""
import os
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fixtures")
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "tools", "pylibs"))


def build_with_openpyxl():
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据表"

    header = ["姓名", "部门", "金额", "日期", "等级", "备注"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    rows = [
        ("张三", "销售部", 128000, "2024-05-12", "A", "重点客户"),
        ("李四", "销售部", 98500, "2024-05-14", "B", ""),
        ("王五", "技术部", 156000, "2024-05-18", "A", "年度框架"),
        ("赵六", "财务部", 64000, "2024-06-02", "C", ""),
        ("钱七", "市场部", 88000, "2024-06-11", "B", "含返点"),
        ("孙八", "技术部", 210000, "2024-06-20", "A", ""),
        ("周九", "销售部", 45000, "2024-07-01", "C", "新客户"),
        ("吴十", "行政部", 32000, "2024-07-05", "C", ""),
    ]
    for name, dept, amount, date, grade, note in rows:
        ws.append([name, dept, amount, date, grade, note])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=6):
        for cell in row:
            cell.border = border
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=3).number_format = "#,##0.00"
        ws.cell(row=row, column=4).number_format = "yyyy-mm-dd"

    ws.merge_cells("A11:C11")
    ws["A11"] = "合计"
    ws["A11"].alignment = Alignment(horizontal="center")
    ws["D11"] = "=SUM(C2:C9)"
    ws["D11"].number_format = "#,##0.00"

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["F"].width = 18
    ws["F2"].comment = Comment("这是批注", "tester")
    ws["A2"].hyperlink = "https://example.com/report"

    ws.conditional_formatting.add(
        "C2:C9", CellIsRule(operator="greaterThan", formula=["100000"], fill=PatternFill("solid", fgColor="FFC7CE"))
    )
    validation = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add("E2:E9")

    ws2 = wb.create_sheet("明细")
    ws2.append(["序号", "项目", "金额"])
    for i in range(1, 31):
        ws2.append([i, f"项目-{i}", i * 137.5])
    ws2["E1"] = "公式测试"
    ws2["E2"] = "=SUM(C2:C31)"

    chart = BarChart()
    chart.title = "金额分布"
    chart.add_data(Reference(ws2, min_col=3, min_row=1, max_row=12), titles_from_data=True)
    ws2.add_chart(chart, "G2")

    ws3 = wb.create_sheet("汇总")
    ws3["A1"] = "汇总页"
    ws3["A2"] = "=明细!E2"
    ws3.sheet_state = "hidden"

    ws4 = wb.create_sheet("已有页眉")
    ws4["A1"] = "这一页原本就有页眉页脚"
    ws4.oddHeader.center.text = "第 &P 页"
    ws4.oddFooter.center.text = "公司内部文件"
    ws4.evenHeader.center.text = "偶数页页眉"
    ws4.HeaderFooter.differentOddEven = True

    path = os.path.join(OUT, "fixture-openpyxl.xlsx")
    wb.save(path)
    print("written", path)


def build_with_xlsxwriter():
    try:
        import xlsxwriter
    except ImportError:
        print("skip xlsxwriter fixture")
        return

    png = os.path.join(HERE, "..", "tools", "_wm.png")
    if not os.path.exists(png):
        print("skip xlsxwriter fixture (missing tools/_wm.png)")
        return

    path = os.path.join(OUT, "fixture-xlsxwriter.xlsx")
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet("带表格")
    ws.write_row("A1", ["列1", "列2"])
    for row in range(1, 12):
        ws.write(row, 0, f"值-{row}")
        ws.write(row, 1, row * 3)
    ws.add_table("A1:B12", {"columns": [{"header": "列1"}, {"header": "列2"}]})
    ws.set_background(png)
    ws.set_header('&L左侧&R右侧')
    ws2 = wb.add_worksheet("普通表")
    ws2.write("A1", "普通内容")
    wb.close()
    print("written", path)


def build_minimal_zip():
    """最小结构：只有必需部件，用来验证健壮性。"""
    path = os.path.join(OUT, "fixture-minimal.xlsx")
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        "</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    workbook = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="唯一表" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    workbook_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        "</Relationships>"
    )
    # 故意不声明 xmlns:r，测试命名空间补全
    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<dimension ref="A1"/><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>hi</t></is></c></row></sheetData>'
        "</worksheet>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", root_rels)
        zf.writestr("xl/workbook.xml", workbook)
        zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        zf.writestr("xl/worksheets/sheet1.xml", sheet)
    print("written", path)


def build_edge_cases():
    """边界结构：sheetPr 里嵌套 extLst、已有 drawing/legacyDrawing/extLst/tableParts。

    用来验证 <picture> 会被插入到 schema 要求的位置（legacyDrawing 之后、extLst 之前），
    并且不会被 sheetPr 内部的 extLst 误导。
    """
    path = os.path.join(OUT, "fixture-edge-cases.xlsx")
    main_ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    r_ns = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    x14_ns = 'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"'

    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
        '<Override PartName="/xl/tables/table1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"/>'
        "</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    workbook = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<workbook {main_ns} {r_ns}>'
        '<sheets><sheet name="边界表" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    workbook_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        "</Relationships>"
    )
    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f"<worksheet {main_ns} {r_ns}>"
        f'<sheetPr><extLst><ext uri="{{X14}}"><x14:foo {x14_ns}/></ext></extLst></sheetPr>'
        '<dimension ref="A1:B3"/>'
        '<sheetViews><sheetView workbookViewId="0"/></sheetViews>'
        '<sheetFormatPr defaultRowHeight="15"/>'
        '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>边界</t></is></c></row></sheetData>'
        '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
        '<drawing r:id="rId1"/>'
        '<legacyDrawing r:id="rId2"/>'
        f'<extLst><ext uri="{{X14}}"><x14:sparklineGroups {x14_ns}/></ext></extLst>'
        '<tableParts count="1"><tablePart r:id="rId3"/></tableParts>'
        "</worksheet>"
    )
    sheet_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" Target="../drawings/vmlDrawing1.vml"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/table" Target="../tables/table1.xml"/>'
        "</Relationships>"
    )
    drawing = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
    )
    table = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<table xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" id="1" name="Table1" displayName="Table1" ref="A1:B3">'
        '<tableColumns count="2"><tableColumn id="1" name="a"/><tableColumn id="2" name="b"/></tableColumns>'
        '<tableStyleInfo name="TableStyleMedium2" showRowStripes="1"/></table>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", root_rels)
        zf.writestr("xl/workbook.xml", workbook)
        zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        zf.writestr("xl/worksheets/sheet1.xml", sheet)
        zf.writestr("xl/worksheets/_rels/sheet1.xml.rels", sheet_rels)
        zf.writestr("xl/drawings/drawing1.xml", drawing)
        zf.writestr("xl/tables/table1.xml", table)
    print("written", path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    build_with_openpyxl()
    build_with_xlsxwriter()
    build_minimal_zip()
    build_edge_cases()
