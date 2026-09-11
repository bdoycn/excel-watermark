#!/usr/bin/env python3
"""用 openpyxl 读取加水印后的文件，输出关键信息（JSON）供 node 测试断言。"""
import json
import sys

import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path)
result = {
    "sheetnames": wb.sheetnames,
    "A2": None,
    "D11": None,
    "hyperlinks": [],
    "headers": {},
}
if "数据表" in wb.sheetnames:
    ws = wb["数据表"]
    result["A2"] = ws["A2"].value
    result["D11"] = ws["D11"].value
    result["merged"] = [str(rng) for rng in ws.merged_cells.ranges]
    result["hyperlinks"] = [cell.hyperlink.target for cell in ws["A"] if cell.hyperlink]
if "明细" in wb.sheetnames:
    result["E2"] = wb["明细"]["E2"].value
    result["charts"] = len(wb["明细"]._charts)
for name in wb.sheetnames:
    ws = wb[name]
    header = ws.oddHeader.center.text if ws.oddHeader and ws.oddHeader.center else None
    footer = ws.oddFooter.center.text if ws.oddFooter and ws.oddFooter.center else None
    even = ws.evenHeader.center.text if ws.evenHeader and ws.evenHeader.center else None
    if header or footer or even:
        result["headers"][name] = {"odd": header, "even": even, "footer": footer}
print(json.dumps(result, ensure_ascii=False))
